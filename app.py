from pathlib import Path
import streamlit as st
import datetime
import json
import copy
import hashlib
import io
from urllib.parse import urlparse
import requests

st.set_page_config(page_title="Ambulance Medication Inventory", page_icon="💊", layout="wide")

# ============================================================
# 1. MASTER MEDICATION DEFINITIONS
# ============================================================
# The Supabase medication_master table is the only source of truth.
# Medication IDs are permanent keys.  The master controls medication
# name, controlled status, expiration tracking, and active status.
# Rig-specific Shelf/Bag par levels live in medication_inventory.
# ============================================================

MEDICATIONS = {}
RIGS = ["Rig #356", "Rig #357"]

# ============================================================
# 2. INVENTORY INITIALIZATION
# ============================================================

def build_empty_inventory():
    return {
        med_id: {
            "shelf_count": 0,
            "bag_count": 0,
            "shelf_par": 0,
            "bag_par": 0,
            "expiry": None,
        }
        for med_id in MEDICATIONS
    }


def build_initial_inventory():
    return {rig: build_empty_inventory() for rig in RIGS}


def purge_inactive_inventory(data):
    """Keep only active medication IDs while preserving historical inventory data."""
    if not isinstance(data, dict):
        return {}
    active_ids = set(MEDICATIONS.keys())
    cleaned = {}
    for rig in RIGS:
        rig_data = data.get(rig, {})
        cleaned[rig] = {
            med_id: record
            for med_id, record in rig_data.items()
            if med_id in active_ids
        } if isinstance(rig_data, dict) else {}
    return cleaned


def normalize_inventory(data):
    """Normalize rig-specific Shelf/Bag counts and par levels without inventing expirations."""
    if not isinstance(data, dict):
        data = {}
    normalized = {}
    for rig in RIGS:
        source = data.get(rig, {})
        normalized[rig] = {}
        for med_id, med in MEDICATIONS.items():
            old = source.get(med_id, {}) if isinstance(source, dict) else {}
            expiry = old.get("expiry")
            if expiry in ("", "None", "null"):
                expiry = None
            normalized[rig][med_id] = {
                "shelf_count": max(0, int(old.get("shelf_count", old.get("count", 0)) or 0)),
                "bag_count": max(0, int(old.get("bag_count", 0) or 0)),
                "shelf_par": max(0, int(old.get("shelf_par", old.get("max", 0)) or 0)),
                "bag_par": max(0, int(old.get("bag_par", 0) or 0)),
                "expiry": str(expiry) if expiry else None,
            }
    return normalized


def initialize_inventory_state():
    if "inventory" not in st.session_state:
        st.session_state.inventory = build_initial_inventory()
    else:
        st.session_state.inventory = normalize_inventory(st.session_state.inventory)

    st.session_state.setdefault("shift_usage", {})
    st.session_state.setdefault("shift_restock", {})
    st.session_state.setdefault("activity_log", [])
    st.session_state.setdefault("par_unlocked", False)
    st.session_state.setdefault("usage_history", [])


# ============================================================
# 3. PERSISTENT USER DATABASE & PERMISSIONS
# ============================================================

SUPABASE_URL = st.secrets.get("supabase", {}).get("url", "").strip().rstrip("/")
SUPABASE_SERVICE_ROLE_KEY = st.secrets.get("supabase", {}).get("service_role_key", "").strip()
SUPABASE_TABLE = "app_users"
SUPABASE_INVENTORY_TABLE = "medication_inventory"
SUPABASE_MEDICATION_TABLE = "medication_master"


def validate_supabase_url(url):
    if not url:
        return "", "Supabase URL is missing from Streamlit Secrets."
    if not url.startswith(("https://", "http://")):
        url = "https://" + url
    parsed = urlparse(url)
    if parsed.scheme != "https" or not parsed.netloc:
        return "", "Supabase URL must look like https://YOUR_PROJECT_REF.supabase.co"
    host = parsed.netloc.lower()
    if not host.endswith(".supabase.co") or any(ch in host for ch in [" ", "=", '"', "'", "<", ">", "\t", "\r", "\n"]):
        return "", "The Supabase URL is invalid. Use the Project URL from Supabase → Settings → API."
    if parsed.path not in ("", "/") or parsed.query or parsed.fragment:
        return "", "Supabase URL should contain only the project URL."
    return url.rstrip("/"), ""


SUPABASE_URL, SUPABASE_URL_ERROR = validate_supabase_url(SUPABASE_URL)

PERMISSION_LABELS = {
    "manage_users": "Manage users",
    "manage_minmax": "Manage medication par levels",
    "edit_inventory": "Edit inventory",
    "record_usage": "Record medication usage",
    "record_restock": "Record medication restock",
}

DEFAULT_ADMIN = {
    "name": "Administrator",
    "initials": "ADM",
    "pin_hash": hashlib.sha256("1234".encode()).hexdigest(),
    "permissions": list(PERMISSION_LABELS.keys()),
    "provider_level": "Paramedic",
    "active": True,
}


def hash_pin(pin):
    return hashlib.sha256(str(pin).encode()).hexdigest()


def supabase_configured():
    return bool(SUPABASE_URL and SUPABASE_SERVICE_ROLE_KEY and not SUPABASE_URL_ERROR)


def supabase_headers():
    return {
        "apikey": SUPABASE_SERVICE_ROLE_KEY,
        "Authorization": f"Bearer {SUPABASE_SERVICE_ROLE_KEY}",
        "Content-Type": "application/json",
    }


def load_users():
    if not supabase_configured():
        return []
    try:
        response = requests.get(
            f"{SUPABASE_URL}/rest/v1/{SUPABASE_TABLE}",
            headers={**supabase_headers(), "Accept": "application/json"},
            params={"select": "name,initials,pin_hash,permissions,provider_level,active", "order": "initials.asc"},
            timeout=10,
        )
        response.raise_for_status()
        rows = response.json()
        if not isinstance(rows, list):
            raise ValueError("Supabase returned an invalid user database.")
        users = [
            {
                "name": row["name"],
                "initials": row["initials"],
                "pin_hash": row["pin_hash"],
                "permissions": row.get("permissions") or [],
                "provider_level": row.get("provider_level") or "EMT / Basic",
                "active": bool(row.get("active", True)),
            }
            for row in rows
        ]
        if not users:
            if save_users([copy.deepcopy(DEFAULT_ADMIN)]):
                return [copy.deepcopy(DEFAULT_ADMIN)]
        return users
    except (requests.RequestException, ValueError, KeyError, TypeError) as exc:
        st.error(f"Unable to load the persistent user database from Supabase. Details: {exc}")
        return []


def save_users(users):
    if not supabase_configured():
        st.error("Cannot save users because the Supabase database is not configured.")
        return False
    rows = [
        {
            "name": user["name"],
            "initials": user["initials"].upper(),
            "pin_hash": user["pin_hash"],
            "permissions": user.get("permissions", []),
            "provider_level": user.get("provider_level", "EMT / Basic"),
            "active": bool(user.get("active", True)),
        }
        for user in users
    ]
    try:
        response = requests.post(
            f"{SUPABASE_URL}/rest/v1/{SUPABASE_TABLE}",
            headers={**supabase_headers(), "Prefer": "resolution=merge-duplicates,return=minimal"},
            json=rows,
            timeout=10,
        )
        response.raise_for_status()
        return True
    except requests.RequestException as exc:
        st.error(f"Unable to save the user database to Supabase. Details: {exc}")
        return False



def medication_master_rows():
    return [
        {
            "medication_id": med_id,
            "name": med["name"],
            "controlled": bool(med["controlled"]),
            "track_expiration": bool(med.get("track_expiration", False)),
            "active": True,
        }
        for med_id, med in MEDICATIONS.items()
    ]


def apply_medication_master(rows):
    """Replace the in-memory medication master with validated active Supabase rows."""
    global MEDICATIONS
    new_meds = {}
    for row in rows:
        med_id = str(row["medication_id"]).strip()
        new_meds[med_id] = {
            "id": med_id,
            "name": str(row["name"]).strip(),
            "controlled": bool(row.get("controlled", False)),
            "track_expiration": bool(row.get("track_expiration", False)),
        }
    MEDICATIONS = new_meds


def load_medication_master(force=False):
    """Load the medication master strictly from Supabase."""
    if not supabase_configured():
        st.session_state.medication_master_load_error = "Supabase is not configured."
        return False
    if st.session_state.get("medication_master_loaded") and not force:
        return bool(MEDICATIONS)
    try:
        response = requests.get(
            f"{SUPABASE_URL}/rest/v1/{SUPABASE_MEDICATION_TABLE}",
            headers={**supabase_headers(), "Accept": "application/json"},
            params={
                "select": "medication_id,name,controlled,track_expiration,active",
                "order": "medication_id.asc",
            },
            timeout=10,
        )
        response.raise_for_status()
        rows = response.json()
        if not isinstance(rows, list):
            raise ValueError("Supabase returned an invalid medication master list.")

        validated = []
        seen = set()
        for row in rows:
            if not bool(row.get("active", True)):
                continue
            med_id = str(row.get("medication_id", "")).strip()
            name = str(row.get("name", "")).strip()
            if not med_id or not name or med_id in seen:
                continue
            validated.append({
                "medication_id": med_id,
                "name": name,
                "controlled": bool(row.get("controlled", False)),
                "track_expiration": bool(row.get("track_expiration", False)),
                "active": True,
            })
            seen.add(med_id)

        apply_medication_master(validated)
        st.session_state.medication_master_loaded = True
        st.session_state.medication_master_load_error = ""
        return bool(validated)
    except (requests.RequestException, ValueError, TypeError) as exc:
        st.session_state.medication_master_load_error = str(exc)
        return False


def save_medication_master_rows(rows):
    """Upsert the complete medication master and deactivate omitted IDs."""
    if not supabase_configured():
        st.error("Cannot save the medication master list because the Supabase database is not configured.")
        return False
    try:
        existing_response = requests.get(
            f"{SUPABASE_URL}/rest/v1/{SUPABASE_MEDICATION_TABLE}",
            headers={**supabase_headers(), "Accept": "application/json"},
            params={"select": "medication_id,name,controlled,track_expiration,active"},
            timeout=10,
        )
        existing_response.raise_for_status()
        existing_rows = existing_response.json()
        if not isinstance(existing_rows, list):
            raise ValueError("Supabase returned an invalid existing medication master.")

        uploaded_ids = {str(row["medication_id"]).strip() for row in rows}
        merged = []
        for row in rows:
            merged.append({
                "medication_id": str(row["medication_id"]).strip(),
                "name": str(row["name"]).strip(),
                "controlled": bool(row["controlled"]),
                "track_expiration": bool(row.get("track_expiration", False)),
                "active": bool(row["active"]),
            })
        for old in existing_rows:
            old_id = str(old.get("medication_id", "")).strip()
            if old_id and old_id not in uploaded_ids:
                merged.append({
                    "medication_id": old_id,
                    "name": str(old.get("name", old_id)),
                    "controlled": bool(old.get("controlled", False)),
                    "track_expiration": bool(old.get("track_expiration", False)),
                    "active": False,
                })

        response = requests.post(
            f"{SUPABASE_URL}/rest/v1/{SUPABASE_MEDICATION_TABLE}",
            headers={**supabase_headers(), "Prefer": "resolution=merge-duplicates,return=representation"},
            json=merged,
            timeout=15,
        )
        response.raise_for_status()
        return True
    except (requests.RequestException, ValueError, TypeError) as exc:
        st.error(f"Unable to save the medication master list to Supabase. Details: {exc}")
        return False


def medication_master_excel_bytes():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "Medications"
    headers = ["Medication ID", "Medication Name", "Controlled", "Track Expiration", "Active"]
    ws.append(headers)
    for row in medication_master_rows():
        ws.append([
            row["medication_id"], row["name"], row["controlled"],
            row["track_expiration"], row["active"],
        ])

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    widths = {"A": 28, "B": 42, "C": 14, "D": 20, "E": 10}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    info = wb.create_sheet("Instructions")
    instructions = [
        ["Ambulance Medication Master List"],
        ["Edit the Medications sheet, then upload it back into the app."],
        ["Medication ID is the permanent key. Do not change an existing ID if you want to preserve inventory."],
        ["Medication Name can be corrected or updated."],
        ["Controlled: TRUE hides the medication from EMT / Basic users; FALSE makes it visible."],
        ["Track Expiration: TRUE requires expiration tracking. Enter expiration as MM/YY in the app; the system stores the last day of that month."],
        ["Active: TRUE keeps the medication in the active list. FALSE removes it from the active list without deleting its historical inventory rows."],
        ["Shelf/Bag par levels are configured separately for each rig and are not part of this master list."],
        ["To add a medication, use a new unique Medication ID."],
    ]
    for row in instructions:
        info.append(row)
    info.column_dimensions["A"].width = 120
    info["A1"].font = Font(bold=True, size=14)
    info.freeze_panes = "A2"

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def validate_medication_master_upload(uploaded_file):
    from openpyxl import load_workbook

    if uploaded_file is None:
        return [], ["No spreadsheet was selected."]
    try:
        wb = load_workbook(uploaded_file, read_only=True, data_only=True)
        if "Medications" not in wb.sheetnames:
            return [], ["The spreadsheet must contain a sheet named 'Medications'."]
        ws = wb["Medications"]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return [], ["The Medications sheet is empty."]

        headers = [str(v).strip() if v is not None else "" for v in rows[0]]
        required = ["Medication ID", "Medication Name", "Controlled", "Track Expiration", "Active"]
        if headers[:len(required)] != required:
            return [], ["The first five columns must be: Medication ID, Medication Name, Controlled, Track Expiration, Active."]

        errors, cleaned, seen = [], [], set()

        def parse_bool(value, field, excel_row):
            if isinstance(value, bool):
                return value
            text = str(value).strip().lower()
            if text in {"true", "yes", "y", "1"}:
                return True
            if text in {"false", "no", "n", "0"}:
                return False
            errors.append(f"Row {excel_row}: {field} must be TRUE or FALSE.")
            return False

        for excel_row, values in enumerate(rows[1:], start=2):
            if all(v is None or str(v).strip() == "" for v in values):
                continue
            med_id = str(values[0]).strip() if len(values) > 0 and values[0] is not None else ""
            name = str(values[1]).strip() if len(values) > 1 and values[1] is not None else ""
            controlled_raw = values[2] if len(values) > 2 else None
            expiration_raw = values[3] if len(values) > 3 else None
            active_raw = values[4] if len(values) > 4 else None

            if not med_id:
                errors.append(f"Row {excel_row}: Medication ID is required.")
            elif med_id in seen:
                errors.append(f"Row {excel_row}: Duplicate Medication ID '{med_id}'.")
            if not name:
                errors.append(f"Row {excel_row}: Medication Name is required.")

            controlled = parse_bool(controlled_raw, "Controlled", excel_row)
            track_expiration = parse_bool(expiration_raw, "Track Expiration", excel_row)
            active = parse_bool(active_raw, "Active", excel_row)

            if med_id:
                seen.add(med_id)
            cleaned.append({
                "medication_id": med_id,
                "name": name,
                "controlled": controlled,
                "track_expiration": track_expiration,
                "active": active,
            })

        if not cleaned:
            errors.append("No medication rows were found.")
        return cleaned, errors
    except Exception as exc:
        return [], [f"Could not read the spreadsheet: {exc}"]


def inventory_rows_from_state():
    rows = []
    for rig in RIGS:
        for med_id in MEDICATIONS:
            item = st.session_state.inventory[rig][med_id]
            rows.append({
                "rig": rig,
                "medication_id": med_id,
                "shelf_count": int(item.get("shelf_count", 0)),
                "bag_count": int(item.get("bag_count", 0)),
                "shelf_par": int(item.get("shelf_par", 0)),
                "bag_par": int(item.get("bag_par", 0)),
                "expiry": item.get("expiry"),
            })
    return rows


def save_inventory_rows(rows):
    if not supabase_configured():
        st.error("Cannot save inventory because the Supabase database is not configured.")
        return False
    try:
        response = requests.post(
            f"{SUPABASE_URL}/rest/v1/{SUPABASE_INVENTORY_TABLE}",
            headers={**supabase_headers(), "Prefer": "resolution=merge-duplicates,return=minimal"},
            json=rows,
            timeout=10,
        )
        response.raise_for_status()
        return True
    except requests.RequestException as exc:
        st.error(f"Unable to save inventory to Supabase. Details: {exc}")
        return False


def save_inventory_item(rig, med_id):
    item = st.session_state.inventory[rig][med_id]
    row = {
        "rig": rig,
        "medication_id": med_id,
        "shelf_count": int(item["shelf_count"]),
        "bag_count": int(item["bag_count"]),
        "shelf_par": int(item["shelf_par"]),
        "bag_par": int(item["bag_par"]),
        "expiry": item.get("expiry"),
    }
    return save_inventory_rows([row])


def initialize_inventory_from_supabase():
    initialize_inventory_state()
    if not supabase_configured() or st.session_state.get("inventory_loaded_from_supabase"):
        return True
    try:
        response = requests.get(
            f"{SUPABASE_URL}/rest/v1/{SUPABASE_INVENTORY_TABLE}",
            headers={**supabase_headers(), "Accept": "application/json"},
            params={
                "select": "rig,medication_id,shelf_count,bag_count,shelf_par,bag_par,expiry",
            },
            timeout=10,
        )
        response.raise_for_status()
        rows = response.json()
        if not isinstance(rows, list):
            raise ValueError("Supabase returned an invalid inventory database.")

        for row in rows:
            rig = row.get("rig")
            med_id = row.get("medication_id")
            if rig in RIGS and med_id in MEDICATIONS:
                expiry = row.get("expiry")
                if expiry in ("", "None", "null"):
                    expiry = None
                st.session_state.inventory[rig][med_id] = {
                    "shelf_count": max(0, int(row.get("shelf_count", 0) or 0)),
                    "bag_count": max(0, int(row.get("bag_count", 0) or 0)),
                    "shelf_par": max(0, int(row.get("shelf_par", 0) or 0)),
                    "bag_par": max(0, int(row.get("bag_par", 0) or 0)),
                    "expiry": str(expiry) if expiry else None,
                }

        st.session_state.inventory = normalize_inventory(st.session_state.inventory)
        if not rows:
            save_inventory_rows(inventory_rows_from_state())
        st.session_state.inventory_loaded_from_supabase = True
        return True
    except (requests.RequestException, ValueError, TypeError) as exc:
        st.error(f"Unable to load medication inventory from Supabase. Details: {exc}")
        return False


def save_usage_history(rig, med_id, location, quantity):
    if not supabase_configured():
        st.error("Cannot save usage history because the Supabase database is not configured.")
        return False
    row = {
        "rig": rig,
        "medication_id": med_id,
        "location": location,
        "quantity": int(quantity),
        "user_initials": (current_user() or {}).get("initials", "SYSTEM"),
    }
    try:
        response = requests.post(
            f"{SUPABASE_URL}/rest/v1/medication_usage",
            headers={**supabase_headers(), "Prefer": "return=minimal"},
            json=[row],
            timeout=10,
        )
        response.raise_for_status()
        return True
    except requests.RequestException as exc:
        st.error(f"Unable to save medication usage history. Details: {exc}")
        return False


def load_usage_history(rig):
    if not supabase_configured():
        return []
    try:
        response = requests.get(
            f"{SUPABASE_URL}/rest/v1/medication_usage",
            headers={**supabase_headers(), "Accept": "application/json"},
            params={
                "select": "id,used_at,rig,medication_id,location,quantity,user_initials",
                "rig": f"eq.{rig}",
                "order": "used_at.desc",
                "limit": "100",
            },
            timeout=10,
        )
        response.raise_for_status()
        rows = response.json()
        if not isinstance(rows, list):
            raise ValueError("Supabase returned an invalid usage history.")
        return rows
    except (requests.RequestException, ValueError, TypeError) as exc:
        st.session_state.usage_history_load_error = str(exc)
        return []


def current_user():
    return st.session_state.get("current_user")


def has_permission(permission):
    user = current_user()
    return bool(user and user.get("active", False) and permission in user.get("permissions", []))


def is_mobile_device():
    """Best-effort device detection so phones get a purpose-built compact UI."""
    try:
        user_agent = st.context.headers.get("User-Agent", "").lower()
        mobile_terms = (
            "android", "iphone", "ipad", "ipod", "mobile",
            "windows phone", "opera mini", "iemobile"
        )
        return any(term in user_agent for term in mobile_terms)
    except Exception:
        return False


if "users" not in st.session_state:
    st.session_state.users = load_users()

st.session_state.setdefault("current_user", None)
st.session_state.setdefault("minmax_unlocked", False)


# ============================================================
# LOGIN MUST HAPPEN BEFORE MEDICATION MASTER VALIDATION
# ============================================================

if not current_user():
    with st.sidebar.expander("👤 User Access", expanded=True):
        login_initials = st.text_input(
            "Initials",
            max_chars=5,
            key="login_initials"
        ).strip().upper()

        login_pin = st.text_input(
            "PIN",
            type="password",
            key="login_pin"
        )

        if st.button("🔓 Sign In", key="sign_in"):
            match = next(
                (
                    u for u in st.session_state.users
                    if u["active"]
                    and u["initials"].upper() == login_initials
                    and u["pin_hash"] == hash_pin(login_pin)
                ),
                None,
            )

            if match:
                st.session_state.current_user = copy.deepcopy(match)
                st.session_state.minmax_unlocked = False
                st.rerun()
            else:
                st.error("Invalid initials or PIN.")

    st.stop()


# ============================================================
# LOAD MEDICATION MASTER AFTER SUCCESSFUL LOGIN
# ============================================================

if not load_medication_master():
    st.error(
        "The medication master could not be loaded from Supabase. "
        "The inventory cannot be started safely."
    )

    if st.session_state.get("medication_master_load_error"):
        st.caption(
            st.session_state["medication_master_load_error"]
        )

    st.stop()

if not MEDICATIONS:
    st.error(
        "No active medications are defined in the Supabase "
        "medication master. Upload and apply the medication "
        "master list before using inventory."
    )
    st.stop()

initialize_inventory_from_supabase()

# ============================================================
# 4. HELPERS
# ============================================================

def visible_med_ids(rig, role):
    del rig
    return [
        med_id for med_id, med in MEDICATIONS.items()
        if role == "Paramedic" or not med["controlled"]
    ]


def build_visible_medication_list(rig, role):
    return [
        {
            "id": med_id,
            "name": MEDICATIONS[med_id]["name"],
            "shelf_count": st.session_state.inventory[rig][med_id]["shelf_count"],
            "bag_count": st.session_state.inventory[rig][med_id]["bag_count"],
            "shelf_par": st.session_state.inventory[rig][med_id]["shelf_par"],
            "bag_par": st.session_state.inventory[rig][med_id]["bag_par"],
            "expiry": st.session_state.inventory[rig][med_id]["expiry"],
            "track_expiration": MEDICATIONS[med_id].get("track_expiration", False),
        }
        for med_id in visible_med_ids(rig, role)
    ]


def record_activity(rig, med_id, action, quantity, location=None, expiration=None):
    event = {
        "timestamp": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "user": (current_user() or {}).get("initials", "SYSTEM"),
        "rig": rig,
        "medication": MEDICATIONS[med_id]["name"],
        "medication_id": med_id,
        "action": action,
        "quantity": int(quantity),
    }
    if location:
        event["location"] = location
    if expiration:
        event["expiration"] = expiration
    st.session_state.activity_log.append(event)


def total_current(item):
    return int(item.get("shelf_count", 0)) + int(item.get("bag_count", 0))


def total_par(item):
    return int(item.get("shelf_par", 0)) + int(item.get("bag_par", 0))


def location_status(count, par):
    count, par = int(count), int(par)
    if par <= 0:
        return "PAR NOT SET"
    if count == 0:
        return "OUT OF STOCK"
    if count < par:
        return "BELOW PAR"
    return "OK"


def get_status(item):
    shelf_status = location_status(item.get("shelf_count", 0), item.get("shelf_par", 0))
    bag_status = location_status(item.get("bag_count", 0), item.get("bag_par", 0))
    if shelf_status == "OUT OF STOCK" and int(item.get("shelf_par", 0)) > 0:
        return "SHELF EMPTY"
    if bag_status == "OUT OF STOCK" and int(item.get("bag_par", 0)) > 0:
        return "BAG EMPTY"
    if shelf_status == "BELOW PAR" or bag_status == "BELOW PAR":
        return "BELOW PAR"
    if shelf_status in {"PAR NOT SET"} or bag_status in {"PAR NOT SET"}:
        return "PAR NOT SET"
    return "OK"


def shift_totals(rig, med_id, action, location=None):
    store = st.session_state.shift_usage if action == "usage" else st.session_state.shift_restock
    return store.get((rig, med_id, location), 0) if location else store.get((rig, med_id), 0)


def build_shift_summary(rig, role):
    inventory = st.session_state.inventory[rig]
    ids = visible_med_ids(rig, role)
    lines = [
        f"AMBULANCE MEDICATION SHIFT SUMMARY — {rig}",
        f"Generated: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"Certification view: {role}",
        "",
        "CURRENT INVENTORY",
        "Medication | Shelf | Shelf Par | Bag | Bag Par | Total | Total Par | Expiration | Status",
        "-" * 135,
    ]
    for med_id in ids:
        item = inventory[med_id]
        expiry = item.get("expiry") or "N/A"
        lines.append(
            f"{MEDICATIONS[med_id]['name']} | {item['shelf_count']} | {item['shelf_par']} | "
            f"{item['bag_count']} | {item['bag_par']} | {total_current(item)} | {total_par(item)} | "
            f"{expiry} | {get_status(item)}"
        )
    lines.extend(["", "SHIFT ACTIVITY"])
    activity = [e for e in st.session_state.activity_log if e["rig"] == rig]
    if activity:
        for event in reversed(activity[-50:]):
            location = f" | {event.get('location')}" if event.get("location") else ""
            lines.append(
                f"{event['timestamp']} | {event['user']} | {event['action'].upper()} | "
                f"{event['medication']} | {event['quantity']}{location}"
            )
    else:
        lines.append("No usage or restock activity recorded.")
    return "\n".join(lines)


def parse_date(value):
    try:
        return datetime.datetime.strptime(str(value), "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return None


def parse_expiration_mm_yy(value):
    """Convert MM/YY to the last calendar day of that month."""
    text = str(value or "").strip()
    try:
        month_text, year_text = text.split("/")
        month = int(month_text)
        year_short = int(year_text)
        if not 1 <= month <= 12 or not 0 <= year_short <= 99:
            return None
        year = 2000 + year_short
        if month == 12:
            next_month = datetime.date(year + 1, 1, 1)
        else:
            next_month = datetime.date(year, month + 1, 1)
        return next_month - datetime.timedelta(days=1)
    except (TypeError, ValueError):
        return None


def format_expiration(expiry):
    date_value = parse_date(expiry)
    return date_value.strftime("%m/%y") if date_value else ""


def expiration_for_display(med_id, expiry):
    if not MEDICATIONS[med_id].get("track_expiration", False):
        return "Not tracked"
    return format_expiration(expiry) or "Not entered"


def save_usage_and_inventory(rig, med_id, location, quantity, old_item):
    """Persist inventory and usage history, rolling inventory back if history fails."""
    if not save_inventory_item(rig, med_id):
        return False
    if save_usage_history(rig, med_id, location, quantity):
        return True

    st.session_state.inventory[rig][med_id] = copy.deepcopy(old_item)
    save_inventory_item(rig, med_id)
    return False


# ============================================================
# 5. LOGIN / USER ACCESS
# ============================================================
# Authentication is intentionally the first screen.  This is especially
# useful on phones, where the inventory should not be visible until the
# user has signed in.
# ============================================================

if not current_user():
    st.markdown("<div style='text-align:center; padding-top:2rem;'>", unsafe_allow_html=True)
    st.title("💊 Ambulance Medication Inventory")
    st.caption("Authorized Personnel Only")
    st.markdown("</div>", unsafe_allow_html=True)

    _, login_col, _ = st.columns([1, 2, 1])
    with login_col:
        with st.container(border=True):
            st.subheader("🔐 Sign In")
            login_initials = st.text_input(
                "Initials", max_chars=5, key="login_initials",
                placeholder="Enter your initials"
            ).strip().upper()
            login_pin = st.text_input(
                "PIN", type="password", key="login_pin",
                placeholder="Enter your PIN"
            )
            if st.button("🔓 Sign In", type="primary", use_container_width=True, key="sign_in"):
                match = next(
                    (
                        u for u in st.session_state.users
                        if u.get("active")
                        and u.get("initials", "").upper() == login_initials
                        and u.get("pin_hash") == hash_pin(login_pin)
                    ),
                    None,
                )
                if match:
                    st.session_state.current_user = copy.deepcopy(match)
                    st.session_state.par_unlocked = False
                    st.rerun()
                else:
                    st.error("Invalid initials or PIN.")
        st.caption("Authorized user accounts are managed by an administrator.")
    st.stop()

# ============================================================
# AUTHENTICATED USER MANAGEMENT
# ============================================================
with st.sidebar.expander("👤 User Access", expanded=True):
    user = current_user()
    st.success(f"Signed in: {user['name']} ({user['initials']})")
    if st.button("🔒 Sign Out", key="sign_out", use_container_width=True):
        st.session_state.current_user = None
        st.session_state.par_unlocked = False
        st.rerun()

    if has_permission("manage_users"):
        st.divider()
        st.subheader("User Management")
        with st.form("add_user_form", clear_on_submit=True):
            new_name = st.text_input("Name")
            new_initials = st.text_input("Initials", max_chars=5).strip().upper()
            new_pin = st.text_input("PIN", type="password")
            new_pin_confirm = st.text_input("Confirm PIN", type="password")
            new_provider_level = st.selectbox("Provider Level", ["EMT / Basic", "Paramedic"])
            new_permissions = st.multiselect(
                "Permissions",
                list(PERMISSION_LABELS.keys()),
                format_func=lambda p: PERMISSION_LABELS[p],
                default=["record_usage", "record_restock"],
            )
            if st.form_submit_button("➕ Add Authorized User"):
                initials_used = {u["initials"].upper() for u in st.session_state.users}
                if not new_name.strip() or not new_initials:
                    st.error("Name and initials are required.")
                elif new_initials in initials_used:
                    st.error("Those initials are already in use.")
                elif len(new_pin) < 4:
                    st.error("PIN must be at least 4 characters.")
                elif new_pin != new_pin_confirm:
                    st.error("PIN confirmation does not match.")
                else:
                    st.session_state.users.append({
                        "name": new_name.strip(),
                        "initials": new_initials,
                        "pin_hash": hash_pin(new_pin),
                        "permissions": new_permissions,
                        "provider_level": new_provider_level,
                        "active": True,
                    })
                    if save_users(st.session_state.users):
                        st.success("Authorized user added.")
                        st.rerun()

        for i, managed_user in enumerate(st.session_state.users):
            with st.container(border=True):
                st.write(f"**{managed_user['name']}** ({managed_user['initials']})")
                st.caption(f"Provider Level: {managed_user.get('provider_level', 'EMT / Basic')}")
                st.caption(
                    ", ".join(
                        PERMISSION_LABELS[p]
                        for p in managed_user.get("permissions", [])
                        if p in PERMISSION_LABELS
                    ) or "No permissions"
                )
                left, middle, right = st.columns(3)
                is_self = user["initials"] == managed_user["initials"]
                with left:
                    if st.button(
                        "Disable" if managed_user["active"] else "Enable",
                        key=f"toggle_user_{i}", disabled=is_self,
                    ):
                        st.session_state.users[i]["active"] = not managed_user["active"]
                        save_users(st.session_state.users)
                        st.rerun()
                with middle:
                    if st.button("Edit Access", key=f"edit_access_button_{i}"):
                        st.session_state[f"edit_access_open_{i}"] = not st.session_state.get(
                            f"edit_access_open_{i}", False
                        )
                with right:
                    if st.button("Reset PIN", key=f"reset_pin_button_{i}"):
                        st.session_state[f"reset_pin_open_{i}"] = True

                if st.session_state.get(f"edit_access_open_{i}", False):
                    with st.form(f"edit_access_form_{i}"):
                        edited_provider_level = st.selectbox(
                            "Provider Level",
                            ["EMT / Basic", "Paramedic"],
                            index=0 if managed_user.get("provider_level", "EMT / Basic") == "EMT / Basic" else 1,
                            key=f"edit_provider_level_{i}",
                        )
                        current_permissions = [
                            p for p in managed_user.get("permissions", [])
                            if p in PERMISSION_LABELS
                        ]
                        edited_permissions = st.multiselect(
                            "Permissions",
                            list(PERMISSION_LABELS.keys()),
                            format_func=lambda p: PERMISSION_LABELS[p],
                            default=current_permissions,
                            key=f"edit_permissions_{i}",
                        )
                        if is_self:
                            st.caption(
                                "Your own Manage users permission cannot be removed here. "
                                "This prevents an administrator from accidentally locking themselves out."
                            )
                        if st.form_submit_button("💾 Save Access Changes"):
                            if is_self and "manage_users" not in edited_permissions:
                                st.error("Keep Manage users enabled for your own account.")
                            else:
                                st.session_state.users[i]["permissions"] = edited_permissions
                                st.session_state.users[i]["provider_level"] = edited_provider_level
                                if is_self:
                                    st.session_state.current_user = copy.deepcopy(st.session_state.users[i])
                                if save_users(st.session_state.users):
                                    st.success(f"Access updated for {managed_user['name']}.")
                                    st.session_state.pop(f"edit_access_open_{i}", None)
                                    st.rerun()

                if st.session_state.get(f"reset_pin_open_{i}", False):
                    with st.form(f"reset_pin_form_{i}"):
                        p1 = st.text_input("New PIN", type="password")
                        p2 = st.text_input("Confirm New PIN", type="password")
                        if st.form_submit_button("Save New PIN"):
                            if len(p1) < 4:
                                st.error("PIN must be at least 4 characters.")
                            elif p1 != p2:
                                st.error("PIN confirmation does not match.")
                            else:
                                st.session_state.users[i]["pin_hash"] = hash_pin(p1)
                                save_users(st.session_state.users)
                                st.session_state.pop(f"reset_pin_open_{i}", None)
                                st.success("PIN reset.")
                                st.rerun()

# ============================================================
# 6. MEDICATION MASTER LIST MANAGEMENT
# ============================================================
if has_permission("manage_minmax"):
    with st.sidebar.expander("💊 Medication Master List", expanded=False):
        st.caption("Export the current medication list, edit it in Excel, then upload it back here.")
        st.caption("Shelf/Bag par levels are configured per rig in the inventory grid.")
        if supabase_configured():
            st.caption("🟢 Supabase connection configured.")
        else:
            st.error("🔴 Supabase is not configured. The medication master cannot be saved.")
        st.download_button(
            "⬇️ Download Medication List",
            data=medication_master_excel_bytes(),
            file_name="ambulance_medication_master.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
            key="download_medication_master",
        )
        uploaded_master = st.file_uploader(
            "Upload revised medication list",
            type=["xlsx"],
            key="medication_master_upload",
            help="Use the Medications sheet from the exported workbook. Medication ID is the permanent key.",
        )
        if uploaded_master is not None:
            upload_rows, upload_errors = validate_medication_master_upload(uploaded_master)
            if upload_errors:
                st.error("The spreadsheet was not applied. Fix these issues and upload it again:")
                for error in upload_errors[:20]:
                    st.write(f"• {error}")
                if len(upload_errors) > 20:
                    st.caption(f"...and {len(upload_errors) - 20} more errors.")
                st.session_state.pop("validated_medication_master", None)
            else:
                current_ids = set(MEDICATIONS.keys())
                uploaded_active_ids = {row["medication_id"] for row in upload_rows if row["active"]}
                removed_ids = sorted(current_ids - uploaded_active_ids)
                new_ids = sorted(uploaded_active_ids - current_ids)
                st.session_state.validated_medication_master = upload_rows
                st.success(f"Spreadsheet validated: {len(upload_rows)} medication rows.")
                if new_ids:
                    st.info(f"New medications: {len(new_ids)}")
                if removed_ids:
                    st.warning(
                        f"Medications being removed from the active list: {len(removed_ids)}. "
                        "Their existing inventory rows will be preserved."
                    )

        validated_rows = st.session_state.get("validated_medication_master")
        if validated_rows:
            st.divider()
            st.write("### Ready to Apply")
            st.caption("The spreadsheet has passed validation. Applying it updates only the medication master; existing Shelf/Bag counts and par levels remain with their medication IDs.")
            with st.form("apply_medication_master_form", clear_on_submit=False):
                apply_clicked = st.form_submit_button(
                    "✅ Apply Medication List", type="primary", use_container_width=True
                )
            if apply_clicked:
                rows_to_apply = st.session_state.get("validated_medication_master") or []
                with st.spinner("Applying medication master list..."):
                    if save_medication_master_rows(rows_to_apply):
                        if load_medication_master(force=True):
                            st.session_state.inventory = purge_inactive_inventory(
                                normalize_inventory(st.session_state.inventory)
                            )
                            inventory_saved = save_inventory_rows(inventory_rows_from_state())
                            if inventory_saved:
                                st.success(
                                    f"✅ Medication master applied. {len(MEDICATIONS)} active medications are now in the system."
                                )
                                st.session_state.pop("validated_medication_master", None)
                                st.rerun()
                            else:
                                st.error("The medication master was saved, but inventory synchronization failed.")
                        else:
                            st.error("The medication master was saved, but it could not be reloaded from Supabase.")


# ============================================================
# 7. OPERATIONS HUB
# ============================================================

mobile_device = is_mobile_device()

st.title("💊 Ambulance Medication Inventory")
st.caption("Rig-specific Shelf/Bag inventory, par levels, usage, restocking, expiration tracking, and shift summaries.")

if mobile_device:
    st.info("📱 Mobile mode — simplified for quick use on a phone.")

with st.sidebar:
    st.header("🛡️ Operations Hub")
    user_role = user.get("provider_level", "EMT / Basic")
    selected_rig = st.radio("Active Ambulance Unit", RIGS)

# ============================================================
# 8. BACKUP UTILITY
# ============================================================
st.sidebar.markdown("---")
st.sidebar.subheader("💾 Backup Utility")
backup_data = {
    "version": 3,
    "created": datetime.datetime.now().isoformat(timespec="seconds"),
    "inventory": st.session_state.inventory,
    "shift_usage": {f"{rig}|{med_id}|{location}": qty for (rig, med_id, location), qty in st.session_state.shift_usage.items()},
    "shift_restock": {f"{rig}|{med_id}|{location}": qty for (rig, med_id, location), qty in st.session_state.shift_restock.items()},
    "activity_log": st.session_state.activity_log,
}
st.sidebar.download_button(
    "⬇️ Download Backup Database",
    data=json.dumps(backup_data, indent=2),
    file_name="ambulance_med_data.json",
    mime="application/json",
)

# ============================================================
# 9. CURRENT VIEW / ALERTS
# ============================================================
raw_inventory = st.session_state.inventory[selected_rig]
visible_meds = build_visible_medication_list(selected_rig, user_role)
today = datetime.date.today()
fifteen_days_out = today + datetime.timedelta(days=15)

empty_meds, below_par_meds, expiring_meds, all_expiration_dates = [], [], [], []
for med in visible_meds:
    item = raw_inventory[med["id"]]
    if med["track_expiration"]:
        exp_date = parse_date(med["expiry"])
        if exp_date:
            all_expiration_dates.append(exp_date)
            if exp_date <= fifteen_days_out:
                expiring_meds.append(f"{med['name']} ({format_expiration(med['expiry'])})")
        elif total_current(item) > 0:
            expiring_meds.append(f"{med['name']} (expiration not entered)")
    if (item["shelf_par"] > 0 and item["shelf_count"] == 0) or (item["bag_par"] > 0 and item["bag_count"] == 0):
        empty_meds.append(med["name"])
    if (item["shelf_par"] > 0 and item["shelf_count"] < item["shelf_par"]) or (item["bag_par"] > 0 and item["bag_count"] < item["bag_par"]):
        below_par_meds.append(med["name"])

st.subheader("⚠️ Inventory & Expiration Status")
col1, col2, col3, col4 = st.columns(4)
col1.metric("Earliest Expiration", str(min(all_expiration_dates)) if all_expiration_dates else "N/A")
col2.metric("🚨 Shelf/Bag Empty", len(empty_meds))
col3.metric("⚠️ Below Par", len(below_par_meds))
col4.metric("⏳ Expiring ≤ 15 Days", len(expiring_meds))
if empty_meds:
    st.error("**CRITICAL — LOCATION EMPTY:** " + ", ".join(empty_meds))
if below_par_meds:
    st.warning("**NOTICE — BELOW PAR:** " + ", ".join(below_par_meds))
if expiring_meds:
    st.info("**NOTICE — EXPIRATION:** " + ", ".join(expiring_meds))

# ============================================================
# 10. PAR ADMIN LOCK + INVENTORY VIEW
# ============================================================
st.divider()
if has_permission("manage_minmax"):
    st.sidebar.divider()
    st.sidebar.subheader("🔐 Shelf/Bag Par Settings")
    if st.session_state.par_unlocked:
        st.sidebar.success("Shelf/Bag par editing is UNLOCKED.")
        if st.sidebar.button("🔒 Lock Par Levels", key="lock_par"):
            st.session_state.par_unlocked = False
            st.rerun()
    elif st.sidebar.button("🔓 Unlock Par Levels", key="unlock_par"):
        st.session_state.par_unlocked = True
        st.rerun()

if mobile_device:
    st.subheader(f"📱 Quick Inventory — {selected_rig}")
    st.caption("Each medication has separate Shelf and Bag quantities and par levels.")

    mobile_med_id = st.selectbox(
        "Medication", [m["id"] for m in visible_meds],
        format_func=lambda x: MEDICATIONS[x]["name"],
        key=f"mobile_med_{selected_rig}",
    )
    mobile_item = raw_inventory[mobile_med_id]
    mobile_status = get_status(mobile_item)

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Shelf", mobile_item["shelf_count"])
    c2.metric("Bag", mobile_item["bag_count"])
    c3.metric("Total", total_current(mobile_item))
    c4.metric("Total Par", total_par(mobile_item))
    expiry_text = expiration_for_display(mobile_med_id, mobile_item.get("expiry"))
    st.info(f"{mobile_status}  •  Expiration: {expiry_text}")

    if has_permission("edit_inventory") or has_permission("manage_minmax"):
        with st.form(f"mobile_inventory_form_{selected_rig}"):
            m1, m2 = st.columns(2)
            with m1:
                mobile_shelf = st.number_input("Shelf Current", min_value=0, step=1, value=int(mobile_item["shelf_count"]), disabled=not has_permission("edit_inventory"))
            with m2:
                mobile_bag = st.number_input("Bag Current", min_value=0, step=1, value=int(mobile_item["bag_count"]), disabled=not has_permission("edit_inventory"))

            if has_permission("manage_minmax") and st.session_state.par_unlocked:
                p1, p2 = st.columns(2)
                with p1:
                    mobile_shelf_par = st.number_input("Shelf Par", min_value=0, step=1, value=int(mobile_item["shelf_par"]))
                with p2:
                    mobile_bag_par = st.number_input("Bag Par", min_value=0, step=1, value=int(mobile_item["bag_par"]))
            else:
                mobile_shelf_par = mobile_item["shelf_par"]
                mobile_bag_par = mobile_item["bag_par"]

            if MEDICATIONS[mobile_med_id].get("track_expiration", False):
                mobile_expiry = st.text_input(
                    "Expiration (MM/YY)", value=format_expiration(mobile_item.get("expiry")),
                    placeholder="MM/YY", disabled=not has_permission("edit_inventory")
                )
            else:
                mobile_expiry = ""
                st.caption("Expiration tracking is not enabled for this medication.")

            if st.form_submit_button("💾 Save Medication", type="primary", use_container_width=True):
                expiry_date = None
                if MEDICATIONS[mobile_med_id].get("track_expiration", False):
                    expiry_date = parse_expiration_mm_yy(mobile_expiry)
                    if expiry_date is None and (mobile_shelf + mobile_bag) > 0:
                        st.error("Enter expiration as MM/YY, such as 10/26.")
                        st.stop()
                if has_permission("manage_minmax") and st.session_state.par_unlocked:
                    mobile_item["shelf_par"] = int(mobile_shelf_par)
                    mobile_item["bag_par"] = int(mobile_bag_par)
                if has_permission("edit_inventory"):
                    mobile_item["shelf_count"] = int(mobile_shelf)
                    mobile_item["bag_count"] = int(mobile_bag)
                    if MEDICATIONS[mobile_med_id].get("track_expiration", False):
                        mobile_item["expiry"] = expiry_date.isoformat() if expiry_date else None
                    else:
                        mobile_item["expiry"] = None
                if save_inventory_item(selected_rig, mobile_med_id):
                    st.success("Medication inventory saved to Supabase.")
                    st.rerun()

    with st.expander("📋 View All Medications", expanded=False):
        for med in visible_meds:
            item = raw_inventory[med["id"]]
            st.markdown(
                f"**{med['name']}**  \n"
                f"Shelf: **{item['shelf_count']} / {item['shelf_par']}**  •  Bag: **{item['bag_count']} / {item['bag_par']}**  \n"
                f"Total: **{total_current(item)} / {total_par(item)}**  •  Expiration: {expiration_for_display(med['id'], item.get('expiry'))}  •  **{get_status(item)}**"
            )
            st.divider()
else:
    st.subheader(f"📊 Active Operations Grid — {selected_rig}")
    st.caption("Current quantities are editable with inventory permission. Shelf/Bag par levels require the separate par-level permission and unlock.")

    edited_data = st.data_editor(
        visible_meds,
        column_config={
            "id": st.column_config.TextColumn("Medication ID", disabled=True),
            "name": st.column_config.TextColumn("Medication Name", disabled=True),
            "shelf_count": st.column_config.NumberColumn("Shelf Current", min_value=0, step=1, disabled=not has_permission("edit_inventory")),
            "bag_count": st.column_config.NumberColumn("Bag Current", min_value=0, step=1, disabled=not has_permission("edit_inventory")),
            "shelf_par": st.column_config.NumberColumn("Shelf Par", min_value=0, step=1, disabled=not (has_permission("manage_minmax") and st.session_state.par_unlocked)),
            "bag_par": st.column_config.NumberColumn("Bag Par", min_value=0, step=1, disabled=not (has_permission("manage_minmax") and st.session_state.par_unlocked)),
            "expiry": st.column_config.TextColumn("Expiration (MM/YY)", disabled=True),
            "track_expiration": st.column_config.CheckboxColumn("Track Exp.", disabled=True),
        },
        column_order=["id", "name", "shelf_count", "shelf_par", "bag_count", "bag_par", "expiry", "track_expiration"],
        hide_index=True, use_container_width=True, key=f"grid_editor_{selected_rig}",
    )

    if st.button("💾 Save Inventory / Par Changes", type="primary", disabled=not (has_permission("edit_inventory") or has_permission("manage_minmax"))):
        errors = []
        for row in edited_data:
            med_id = row["id"]
            item = raw_inventory[med_id]
            if int(row["shelf_count"]) < 0 or int(row["bag_count"]) < 0 or int(row["shelf_par"]) < 0 or int(row["bag_par"]) < 0:
                errors.append(f"{row['name']}: quantities and par levels cannot be negative.")
                continue
            if has_permission("manage_minmax") and st.session_state.par_unlocked:
                item["shelf_par"] = int(row["shelf_par"])
                item["bag_par"] = int(row["bag_par"])
            if has_permission("edit_inventory"):
                item["shelf_count"] = int(row["shelf_count"])
                item["bag_count"] = int(row["bag_count"])
        if errors:
            for error in errors:
                st.error(error)
        elif save_inventory_rows(inventory_rows_from_state()):
            st.success("✅ Inventory and Shelf/Bag par changes saved to Supabase.")
            st.rerun()

# ============================================================
# 11. SEPARATE USAGE / RESTOCK
# ============================================================
st.divider()
available_ids = visible_med_ids(selected_rig, user_role)

def process_usage(med_id, location, quantity):
    item = raw_inventory[med_id]
    count_key = "shelf_count" if location == "Shelf" else "bag_count"
    if int(quantity) > int(item[count_key]):
        st.error(f"Cannot record {quantity}. {location} has only {item[count_key]} on hand.")
        return
    old_item = copy.deepcopy(item)
    item[count_key] -= int(quantity)
    if save_usage_and_inventory(selected_rig, med_id, location, quantity, old_item):
        st.session_state.shift_usage[(selected_rig, med_id, location)] = shift_totals(selected_rig, med_id, "usage", location) + int(quantity)
        record_activity(selected_rig, med_id, "usage", quantity, location=location)
        st.success(f"Recorded {quantity} × {MEDICATIONS[med_id]['name']} used from {location}.")
        st.rerun()


def process_restock(med_id, location, quantity, incoming_expiry=None):
    item = raw_inventory[med_id]
    count_key = "shelf_count" if location == "Shelf" else "bag_count"
    item[count_key] += int(quantity)
    if MEDICATIONS[med_id].get("track_expiration", False) and incoming_expiry:
        old_expiry = parse_date(item.get("expiry"))
        item["expiry"] = min(old_expiry, incoming_expiry).isoformat() if old_expiry else incoming_expiry.isoformat()
    if save_inventory_item(selected_rig, med_id):
        st.session_state.shift_restock[(selected_rig, med_id, location)] = shift_totals(selected_rig, med_id, "restock", location) + int(quantity)
        record_activity(selected_rig, med_id, "restock", quantity, location=location, expiration=incoming_expiry.isoformat() if incoming_expiry else None)
        active_exp = format_expiration(item.get("expiry")) if item.get("expiry") else "not entered"
        st.success(f"Recorded {quantity} × {MEDICATIONS[med_id]['name']} to {location}. Active expiration: {active_exp}.")
        st.rerun()

if mobile_device:
    usage_tab, restock_tab = st.tabs(["💉 Record Usage", "📦 Record Restock"])
    with usage_tab:
        st.caption("Record medication removed/used from a specific location.")
        if available_ids:
            usage_med_id = st.selectbox("Medication Used", available_ids, format_func=lambda x: MEDICATIONS[x]["name"], key="mobile_usage_med")
            usage_location = st.selectbox("Location", ["Shelf", "Bag"], key="mobile_usage_location")
            usage_qty = st.number_input("Quantity Used", min_value=1, step=1, value=1, key="mobile_usage_qty")
            if st.button("➖ Record Usage", disabled=not has_permission("record_usage"), key="mobile_record_usage", use_container_width=True):
                process_usage(usage_med_id, usage_location, usage_qty)
    with restock_tab:
        st.caption("Record medication added to a specific location. The earliest expiration remains active when expiration is tracked.")
        if available_ids:
            restock_med_id = st.selectbox("Medication Restocked", available_ids, format_func=lambda x: MEDICATIONS[x]["name"], key="mobile_restock_med")
            restock_location = st.selectbox("Location", ["Shelf", "Bag"], key="mobile_restock_location")
            restock_qty = st.number_input("Quantity Restocked", min_value=1, step=1, value=1, key="mobile_restock_qty")
            if MEDICATIONS[restock_med_id].get("track_expiration", False):
                restock_expiry_text = st.text_input("Incoming Expiration (MM/YY)", placeholder="MM/YY", key="mobile_restock_expiry")
            else:
                restock_expiry_text = ""
                st.caption("Expiration tracking is not enabled for this medication.")
            if st.button("➕ Record Restock", disabled=not has_permission("record_restock"), key="mobile_record_restock", use_container_width=True):
                incoming = None
                if MEDICATIONS[restock_med_id].get("track_expiration", False):
                    incoming = parse_expiration_mm_yy(restock_expiry_text)
                    if incoming is None:
                        st.error("Enter incoming expiration as MM/YY, such as 10/26.")
                        st.stop()
                process_restock(restock_med_id, restock_location, restock_qty, incoming)
else:
    left, right = st.columns(2)
    with left:
        st.subheader("💉 Medication Usage")
        st.caption("Record medication removed/used from a specific location.")
        if available_ids:
            usage_med_id = st.selectbox("Medication Used", available_ids, format_func=lambda x: MEDICATIONS[x]["name"], key="usage_med")
            usage_location = st.selectbox("Location", ["Shelf", "Bag"], key="usage_location")
            usage_qty = st.number_input("Quantity Used", min_value=1, step=1, value=1, key="usage_qty")
            if st.button("➖ Record Usage", disabled=not has_permission("record_usage"), key="record_usage_button"):
                process_usage(usage_med_id, usage_location, usage_qty)
    with right:
        st.subheader("📦 Medication Restock")
        st.caption("Record medication added to a specific location. The earliest expiration remains active when expiration is tracked.")
        if available_ids:
            restock_med_id = st.selectbox("Medication Restocked", available_ids, format_func=lambda x: MEDICATIONS[x]["name"], key="restock_med")
            restock_location = st.selectbox("Location", ["Shelf", "Bag"], key="restock_location")
            restock_qty = st.number_input("Quantity Restocked", min_value=1, step=1, value=1, key="restock_qty")
            if MEDICATIONS[restock_med_id].get("track_expiration", False):
                restock_expiry_text = st.text_input("Incoming Expiration (MM/YY)", placeholder="MM/YY", key="restock_expiry")
            else:
                restock_expiry_text = ""
                st.caption("Expiration tracking is not enabled for this medication.")
            if st.button("➕ Record Restock", disabled=not has_permission("record_restock"), key="record_restock_button"):
                incoming = None
                if MEDICATIONS[restock_med_id].get("track_expiration", False):
                    incoming = parse_expiration_mm_yy(restock_expiry_text)
                    if incoming is None:
                        st.error("Enter incoming expiration as MM/YY, such as 10/26.")
                        st.stop()
                process_restock(restock_med_id, restock_location, restock_qty, incoming)

# ============================================================
# 12. RESTOCK NEEDS + COPY-PASTE REQUEST
# ============================================================
st.divider()
st.subheader("📦 Restock Needs")
restock_rows = []
for med in visible_meds:
    item = raw_inventory[med["id"]]
    shelf_needed = max(0, item["shelf_par"] - item["shelf_count"])
    bag_needed = max(0, item["bag_par"] - item["bag_count"])
    if shelf_needed or bag_needed:
        restock_rows.append({
            "Medication": med["name"],
            "Shelf Current": item["shelf_count"],
            "Shelf Par": item["shelf_par"],
            "Shelf Needed": shelf_needed,
            "Bag Current": item["bag_count"],
            "Bag Par": item["bag_par"],
            "Bag Needed": bag_needed,
            "Total Needed": shelf_needed + bag_needed,
            "Expiration": expiration_for_display(med["id"], item.get("expiry")),
            "Status": get_status(item),
        })
if restock_rows:
    st.dataframe(restock_rows, hide_index=True, use_container_width=True)
    request_lines = [f"Ambulance medication restock request — {selected_rig}", ""]
    for row in restock_rows:
        if row["Shelf Needed"]:
            request_lines.append(f"{row['Medication']} — Shelf: {row['Shelf Needed']} needed (current {row['Shelf Current']}, par {row['Shelf Par']})")
        if row["Bag Needed"]:
            request_lines.append(f"{row['Medication']} — Bag: {row['Bag Needed']} needed (current {row['Bag Current']}, par {row['Bag Par']})")
    st.text_area("Supervisor Restock Request", "\n".join(request_lines), height=220)
else:
    st.success("✅ All visible medications are at their Shelf and Bag par levels.")

# ============================================================
# 13. SHIFT SUMMARY + USAGE HISTORY
# ============================================================
st.divider()
st.subheader("📋 Shift Summary Text Exporter")
summary_text = build_shift_summary(selected_rig, user_role)
st.text_area("Shift Summary", summary_text, height=350)
st.download_button(
    "⬇️ Download Shift Summary",
    data=summary_text,
    file_name=f"{selected_rig.replace('#', '').replace(' ', '_')}_shift_summary.txt",
    mime="text/plain",
)

with st.expander("📝 View Shift Activity Log"):
    rig_activity = [event for event in st.session_state.activity_log if event["rig"] == selected_rig]
    if rig_activity:
        st.dataframe(list(reversed(rig_activity)), hide_index=True, use_container_width=True)
    else:
        st.info("No usage or restock activity has been recorded for this rig yet.")

st.subheader("💉 Medication Usage History")
usage_history = load_usage_history(selected_rig)
if usage_history:
    usage_display = []
    for row in usage_history:
        med_id = row.get("medication_id")
        usage_display.append({
            "Date/Time": row.get("used_at", ""),
            "Medication": MEDICATIONS.get(med_id, {}).get("name", med_id),
            "Location": row.get("location", ""),
            "Quantity": row.get("quantity", 0),
            "User": row.get("user_initials", ""),
        })
    st.dataframe(usage_display, hide_index=True, use_container_width=True)
else:
    if st.session_state.get("usage_history_load_error"):
        st.caption(f"Usage history unavailable: {st.session_state.usage_history_load_error}")
    else:
        st.info("No medication usage has been recorded for this rig yet.")

